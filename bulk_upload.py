# bulk_upload.py
"""
Bulk SBOM upload module for Dependency-Track.
Supports uploading thousands of SBOMs from folders with tag mapping from Excel files.
"""

import os
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Any
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from rich.console import Console
from rich.progress import Progress, SpinnerColumn, TextColumn, BarColumn, TaskProgressColumn
from rich.table import Table, box

from project import get_or_create_project
from bom import upload_bom


console = Console()


def parse_host_filename(filename: str) -> Optional[Tuple[str, str]]:
    """
    Parse Host SBOM filename to extract account number and IP.
    
    Example filename: 22206351557_ip-10-53-64-89._4df91bb73d15db3b_2026-02-17T08-16-56Z.json
    Returns: (account_number, ip) -> ('22206351557', 'ip-10-53-64-89')
    
    Args:
        filename: The SBOM filename (without path)
        
    Returns:
        Tuple of (account_number, ip) or None if parsing fails
    """
    # Remove file extension
    name = Path(filename).stem
    
    # Split by underscore
    parts = name.split('_')
    
    if len(parts) < 2:
        console.print(f"[yellow]Warning: Cannot parse filename '{filename}' - not enough parts[/yellow]")
        return None
    
    account_number = parts[0]
    
    # Second part should be the IP (starts with 'ip-')
    ip_part = parts[1]
    
    # Clean up the IP part - remove trailing dot if present
    ip_part = ip_part.rstrip('.')
    
    # Validate account number (should be numeric)
    if not account_number.isdigit():
        console.print(f"[yellow]Warning: Account number '{account_number}' is not numeric in '{filename}'[/yellow]")
        return None
    
    return (account_number, ip_part)


def generate_host_project_name(account_number: str, ip: str) -> str:
    """
    Generate project name for Host category.
    
    Format: <accountnumber>_Host_<IP>
    Example: 22206351557_Host_ip-10-53-64-89
    """
    return f"{account_number}_Host_{ip}"


def load_tags_from_excel(excel_path: str) -> Dict[str, List[str]]:
    """
    Load tags from Excel file.
    
    Expected format: First column is AccountNumber, remaining columns are tags.
    Header row is expected (column names will be ignored).
    
    Args:
        excel_path: Path to the Excel file
        
    Returns:
        Dictionary mapping account numbers to list of tags
    """
    tags_map: Dict[str, List[str]] = {}
    
    if not os.path.exists(excel_path):
        console.print(f"[red]Error: Excel file not found: {excel_path}[/red]")
        return tags_map
    
    try:
        workbook = load_workbook(excel_path, read_only=True)
        sheet = workbook.active
        
        rows = list(sheet.iter_rows(values_only=True))
        
        if len(rows) < 2:
            console.print("[yellow]Warning: Excel file has no data rows (only header)[/yellow]")
            return tags_map
        
        # Skip header row, process data rows
        for row in rows[1:]:
            if not row or not row[0]:
                continue
                
            account_number = str(row[0]).strip()
            
            # Collect all non-empty values from remaining columns as tags
            tags = []
            for cell in row[1:]:
                if cell is not None and str(cell).strip():
                    tags.append(str(cell).strip())
            
            tags_map[account_number] = tags
        
        workbook.close()
        console.print(f"[green]Loaded tags for {len(tags_map)} accounts from Excel[/green]")
        
    except Exception as e:
        console.print(f"[red]Error reading Excel file: {e}[/red]")
    
    return tags_map


def export_failed_sboms_to_excel(failed_files: List[Dict], category: str) -> Optional[str]:
    """
    Export failed SBOM uploads to an Excel file with all details.
    
    Args:
        failed_files: List of dicts with failure details
        category: Category name (Host, Image, Functions)
        
    Returns:
        Path to the generated Excel file, or None if export failed
    """
    if not failed_files:
        return None
    
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = f"Failed {category} SBOMs"
        
        # Define headers
        headers = ['Filename', 'Filepath', 'Account Number', 'IP', 'Project Name', 'Tags', 'Error']
        
        # Style for header
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        
        # Write data rows
        for row_idx, record in enumerate(failed_files, 2):
            ws.cell(row=row_idx, column=1, value=record.get('filename', ''))
            ws.cell(row=row_idx, column=2, value=record.get('filepath', ''))
            ws.cell(row=row_idx, column=3, value=record.get('account_number', ''))
            ws.cell(row=row_idx, column=4, value=record.get('ip', ''))
            ws.cell(row=row_idx, column=5, value=record.get('project_name', ''))
            ws.cell(row=row_idx, column=6, value=record.get('tags', ''))
            ws.cell(row=row_idx, column=7, value=record.get('error', ''))
        
        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)  # Cap at 50 chars
            ws.column_dimensions[column].width = adjusted_width
        
        # Generate filename with timestamp
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_filename = f"failed_{category.lower()}_sboms_{timestamp}.xlsx"
        
        wb.save(excel_filename)
        wb.close()
        
        return excel_filename
        
    except Exception as e:
        console.print(f"[red]Error exporting to Excel: {e}[/red]")
        return None


def scan_sbom_folder(folder_path: str, extension: str = ".json") -> List[str]:
    """
    Recursively scan folder for SBOM files.
    
    Args:
        folder_path: Root folder to scan
        extension: File extension to look for (default: .json)
        
    Returns:
        List of full file paths to SBOM files
    """
    sbom_files = []
    
    if not os.path.exists(folder_path):
        console.print(f"[red]Error: Folder not found: {folder_path}[/red]")
        return sbom_files
    
    for root, dirs, files in os.walk(folder_path):
        for file in files:
            if file.lower().endswith(extension):
                sbom_files.append(os.path.join(root, file))
    
    console.print(f"[cyan]Found {len(sbom_files)} SBOM files in {folder_path}[/cyan]")
    return sbom_files


def bulk_upload_host(
    sbom_folder: str,
    tags_excel: str,
    version: str = "1.0",
    dry_run: bool = False,
    limit: int = 0
) -> Tuple[int, int, List[str]]:
    """
    Bulk upload Host category SBOMs to Dependency-Track.
    
    Args:
        sbom_folder: Path to folder containing Host SBOMs (can have subfolders)
        tags_excel: Path to Excel file with account number to tags mapping
        version: Project version to use (default: "1.0")
        dry_run: If True, only simulate the upload without actually uploading
        limit: Maximum number of SBOMs to upload (0 = no limit)
        
    Returns:
        Tuple of (success_count, failed_count, list of failed filenames)
    """
    console.print("\n[bold cyan]═══════════════════════════════════════════════════════════[/bold cyan]")
    console.print("[bold cyan]           Bulk Host SBOM Upload to Dependency-Track        [/bold cyan]")
    console.print("[bold cyan]═══════════════════════════════════════════════════════════[/bold cyan]\n")
    
    # Load tags from Excel
    console.print("[bold]Step 1: Loading tags from Excel...[/bold]")
    tags_map = load_tags_from_excel(tags_excel)
    
    # Scan for SBOM files
    console.print("\n[bold]Step 2: Scanning for SBOM files...[/bold]")
    sbom_files = scan_sbom_folder(sbom_folder)
    
    # Apply limit if specified
    if limit > 0 and len(sbom_files) > limit:
        console.print(f"[yellow]Limiting to first {limit} files (out of {len(sbom_files)})[/yellow]")
        sbom_files = sbom_files[:limit]
    
    if not sbom_files:
        console.print("[red]No SBOM files found. Exiting.[/red]")
        return (0, 0, [])
    
    success_count = 0
    failed_count = 0
    failed_files = []  # List of dicts with detailed info
    failed_filenames = []  # Simple list for backward compatibility
    
    console.print(f"\n[bold]Step 3: {'Simulating' if dry_run else 'Uploading'} SBOMs...[/bold]")
    
    with Progress(
        SpinnerColumn(),
        TextColumn("[progress.description]{task.description}"),
        BarColumn(),
        TaskProgressColumn(),
        console=console
    ) as progress:
        task = progress.add_task("[cyan]Processing SBOMs...", total=len(sbom_files))
        
        for sbom_path in sbom_files:
            filename = os.path.basename(sbom_path)
            
            # Parse filename to extract account number and IP
            parsed = parse_host_filename(filename)
            
            if parsed is None:
                failed_count += 1
                failed_files.append({
                    'filename': filename,
                    'filepath': sbom_path,
                    'account_number': 'PARSE_ERROR',
                    'ip': 'PARSE_ERROR',
                    'project_name': 'PARSE_ERROR',
                    'tags': '',
                    'error': 'Failed to parse filename'
                })
                failed_filenames.append(filename)
                progress.update(task, advance=1)
                continue
            
            account_number, ip = parsed
            project_name = generate_host_project_name(account_number, ip)
            
            # Get tags for this account number
            account_tags = tags_map.get(account_number, [])
            
            # Add account number as a tag
            all_tags = [account_number] + account_tags
            
            if dry_run:
                # Just simulate - print what would happen
                progress.console.print(
                    f"[dim]Would create project: {project_name} (v{version}) "
                    f"with tags: {all_tags}[/dim]"
                )
                success_count += 1
            else:
                try:
                    # Create or get project with tags
                    project = get_or_create_project(
                        name=project_name,
                        version=version,
                        tags=all_tags,
                        allow_existing=True
                    )
                    
                    if project:
                        # Upload the BOM
                        upload_bom(project['uuid'], sbom_path)
                        success_count += 1
                    else:
                        failed_count += 1
                        failed_files.append({
                            'filename': filename,
                            'filepath': sbom_path,
                            'account_number': account_number,
                            'ip': ip,
                            'project_name': project_name,
                            'tags': ', '.join(all_tags),
                            'error': 'Failed to create project'
                        })
                        failed_filenames.append(filename)
                        
                except Exception as e:
                    error_msg = str(e)
                    # Extract useful error message
                    if 'must have only unique items' in error_msg:
                        error_msg = 'Duplicate components in SBOM'
                    console.print(f"[red]Error uploading {filename}: {error_msg}[/red]")
                    failed_count += 1
                    failed_files.append({
                        'filename': filename,
                        'filepath': sbom_path,
                        'account_number': account_number,
                        'ip': ip,
                        'project_name': project_name,
                        'tags': ', '.join(all_tags),
                        'error': error_msg[:200]  # Truncate long errors
                    })
                    failed_filenames.append(filename)
            
            progress.update(task, advance=1)
    
    # Print summary
    console.print("\n[bold]═══════════════════════════════════════════════════════════[/bold]")
    console.print("[bold]                         Summary                            [/bold]")
    console.print("[bold]═══════════════════════════════════════════════════════════[/bold]")
    
    summary_table = Table(show_header=False, box=box.SIMPLE)
    summary_table.add_column("Metric", style="cyan")
    summary_table.add_column("Value", style="green")
    
    summary_table.add_row("Total SBOMs processed", str(len(sbom_files)))
    summary_table.add_row("Successful uploads", str(success_count))
    summary_table.add_row("Failed uploads", str(failed_count))
    
    console.print(summary_table)
    
    if failed_files:
        console.print("\n[yellow]Failed files:[/yellow]")
        for f in failed_files[:10]:  # Show first 10 failed files
            console.print(f"  - {f['filename']}: {f['error']}")
        if len(failed_files) > 10:
            console.print(f"  ... and {len(failed_files) - 10} more")
        
        # Export failed files to Excel
        excel_path = export_failed_sboms_to_excel(failed_files, "Host")
        if excel_path:
            console.print(f"\n[green]Failed SBOMs exported to:[/green] {excel_path}")
    
    return (success_count, failed_count, failed_filenames)


def preview_host_parsing(sbom_folder: str, tags_excel: str, limit: int = 10):
    """
    Preview how Host SBOM filenames will be parsed without uploading.
    Useful for validating the parsing logic before bulk upload.
    
    Args:
        sbom_folder: Path to folder containing Host SBOMs
        tags_excel: Path to Excel file with tags
        limit: Maximum number of files to preview
    """
    console.print("\n[bold cyan]Preview: Host SBOM Parsing[/bold cyan]\n")
    
    # Load tags
    tags_map = load_tags_from_excel(tags_excel)
    
    # Scan files
    sbom_files = scan_sbom_folder(sbom_folder)
    
    if not sbom_files:
        return
    
    # Create preview table
    table = Table(show_header=True, header_style="bold magenta", box=box.ROUNDED)
    table.add_column("Filename", style="dim", max_width=50)
    table.add_column("Account #", style="cyan")
    table.add_column("IP", style="green")
    table.add_column("Project Name", style="yellow")
    table.add_column("Tags", style="blue", max_width=40)
    
    for sbom_path in sbom_files[:limit]:
        filename = os.path.basename(sbom_path)
        parsed = parse_host_filename(filename)
        
        if parsed:
            account_number, ip = parsed
            project_name = generate_host_project_name(account_number, ip)
            account_tags = tags_map.get(account_number, [])
            all_tags = [account_number] + account_tags
            tags_str = ", ".join(all_tags) if all_tags else "-"
            
            table.add_row(
                filename[:50] + "..." if len(filename) > 50 else filename,
                account_number,
                ip,
                project_name,
                tags_str
            )
        else:
            table.add_row(
                filename[:50] + "..." if len(filename) > 50 else filename,
                "[red]PARSE ERROR[/red]",
                "-",
                "-",
                "-"
            )
    
    console.print(table)
    
    if len(sbom_files) > limit:
        console.print(f"\n[dim]Showing {limit} of {len(sbom_files)} files. "
                     f"Use --limit to see more.[/dim]")
