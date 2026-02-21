#!/usr/bin/env python3
"""
Simple script to download all projects from Dependency-Track and export to Excel with detailed information.
"""

import sys
import json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import requests
import urllib3
from project import get_projects
from config import get_api_headers, get_dtrack_url

# Disable SSL warnings for self-signed certificates
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

def get_project_details(project_uuid):
    """Get detailed information for a specific project."""
    try:
        # Get project metrics
        metrics_url = f"{get_dtrack_url()}api/v1/metrics/project/{project_uuid}/current"
        metrics_response = requests.get(metrics_url, headers=get_api_headers(), verify=False)
        metrics = metrics_response.json() if metrics_response.status_code == 200 else {}
        
        # Get project vulnerabilities count
        vuln_url = f"{get_dtrack_url()}api/v1/vulnerability/project/{project_uuid}"
        vuln_response = requests.get(vuln_url, headers=get_api_headers(), verify=False)
        vulnerabilities = vuln_response.json() if vuln_response.status_code == 200 else []
        
        # Count vulnerabilities by severity
        vuln_counts = {
            'CRITICAL': 0,
            'HIGH': 0,
            'MEDIUM': 0,
            'LOW': 0,
            'INFO': 0,
            'UNASSIGNED': 0
        }
        
        for vuln in vulnerabilities:
            severity = vuln.get('vulnerability', {}).get('severity', 'UNASSIGNED')
            vuln_counts[severity] = vuln_counts.get(severity, 0) + 1
        
        return {
            'metrics': metrics,
            'vulnerability_counts': vuln_counts,
            'total_vulnerabilities': len(vulnerabilities)
        }
    except Exception as e:
        print(f"Warning: Could not get details for project {project_uuid}: {e}")
        return {
            'metrics': {},
            'vulnerability_counts': {},
            'total_vulnerabilities': 0
        }

def format_tags(tags):
    """Format tags list to a readable string."""
    if not tags:
        return ""
    return ", ".join([tag.get('name', '') for tag in tags])

def export_projects_to_excel():
    """Export all projects with details to Excel file."""
    print("Fetching projects from Dependency-Track...")
    
    # Get all projects
    projects = get_projects()
    if not projects:
        print("No projects found or unable to fetch projects.")
        return False
    
    print(f"Found {len(projects)} projects. Getting detailed information...")
    
    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Projects Export"
    
    # Define headers
    headers = [
        'Project Name', 'Version', 'UUID', 'Active', 'Classifier',
        'Created Date', 'Last BOM Import', 'Last BOM Import Format',
        'Tags', 'Total Components', 'Vulnerable Components',
        'Critical Vulnerabilities', 'High Vulnerabilities', 'Medium Vulnerabilities',
        'Low Vulnerabilities', 'Info Vulnerabilities', 'Unassigned Vulnerabilities',
        'Total Vulnerabilities', 'Policy Violations', 'License Risk Score',
        'Operational Risk Score', 'Technical Risk Score', 'Business Risk Score',
        'Risk Score', 'Inherited Risk Score'
    ]
    
    # Write headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    # Process each project
    for row, project in enumerate(projects, 2):
        print(f"Processing project {row-1}/{len(projects)}: {project.get('name', 'Unknown')}")
        
        # Get project details
        details = get_project_details(project['uuid'])
        metrics = details['metrics']
        vuln_counts = details['vulnerability_counts']
        
        # Write project data
        data = [
            project.get('name', ''),
            project.get('version', ''),
            project.get('uuid', ''),
            'Yes' if project.get('active', False) else 'No',
            project.get('classifier', ''),
            project.get('created', ''),
            project.get('lastBomImport', ''),
            project.get('lastBomImportFormat', ''),
            format_tags(project.get('tags', [])),
            metrics.get('components', 0),
            metrics.get('vulnerableComponents', 0),
            vuln_counts.get('CRITICAL', 0),
            vuln_counts.get('HIGH', 0),
            vuln_counts.get('MEDIUM', 0),
            vuln_counts.get('LOW', 0),
            vuln_counts.get('INFO', 0),
            vuln_counts.get('UNASSIGNED', 0),
            details['total_vulnerabilities'],
            metrics.get('policyViolations', 0),
            metrics.get('licenseRisk', 0),
            metrics.get('operationalRisk', 0),
            metrics.get('technicalRisk', 0),
            metrics.get('businessRisk', 0),
            metrics.get('riskScore', 0),
            metrics.get('inheritedRiskScore', 0)
        ]
        
        for col, value in enumerate(data, 1):
            ws.cell(row=row, column=col, value=value)
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        
        adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Add borders and formatting
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'), 
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for row in ws.iter_rows(min_row=1, max_row=len(projects)+1):
        for cell in row:
            cell.border = thin_border
    
    # Save the file
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"dependency_track_projects_{timestamp}.xlsx"
    
    try:
        wb.save(filename)
        print(f"\n✓ Successfully exported {len(projects)} projects to: {filename}")
        
        # Print summary
        print(f"\nExport Summary:")
        print(f"  • Total Projects: {len(projects)}")
        print(f"  • Export Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        print(f"  • File Location: {filename}")
        
        return True
        
    except Exception as e:
        print(f"Error saving file: {e}")
        return False

def main():
    """Main function to run the export."""
    print("Dependency-Track Projects Export Tool")
    print("=" * 50)
    
    try:
        success = export_projects_to_excel()
        if success:
            print("\nExport completed successfully!")
        else:
            print("\nExport failed!")
            sys.exit(1)
            
    except KeyboardInterrupt:
        print("\nExport cancelled by user.")
        sys.exit(1)
    except Exception as e:
        print(f"\nUnexpected error: {e}")
        sys.exit(1)

if __name__ == "__main__":
    main()